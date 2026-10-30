"""单元测试：编码（含米数写法）、金蝶解析、MES 报文、多语言契约、迁移校验和、cron、导出防注入、配置。"""

import gzip
import json
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest

from app.core import encoding as codec
from app.core.encoding import DEFAULT_CHARSETS, Spec
from app.core.errors import ErrorCode
from app.core.scheduler import Cron
from app.db.migrate import migrations
from app.middleware.request_id import request_id_of
from app.services import files, k3, mes_client

ROOT = Path(__file__).resolve().parents[1]

# ---------------------------------------------------------------------- SN 编码（5.2.2）

HX = Spec("HX", "A", 10, 6, "0123456789", True)
NOLEN = Spec("AVM", "X", 10, 5, "0123456789", False)
B32 = Spec("Q", "", 32, 4, DEFAULT_CHARSETS[32], False)
#: 前缀 + 流水(4) + 米数前分隔符 + 米数（补足 5 位）+ 米数后分隔符 + 单位 + 后缀
FULL = Spec("光-SN", "#尾", 10, 4, "0123456789", True, 5, "-", "M", True)
UNIT_ONLY = Spec("U", "", 10, 3, "0123456789", False, with_unit=True)


def test_sn_is_prefix_serial_length_unit_suffix():
    assert codec.format_sn(1, HX, "100") == "HX000001100A", "流水紧跟前缀，米数在流水后"
    assert codec.format_sn(1, HX, "10.5") == "HX00000110.5A"
    assert codec.format_sn(1, NOLEN, "100") == "AVM00001X", "规则不带米数时不写米数"
    assert codec.format_sn(1, NOLEN, "100", "PCS") == "AVM00001X", "规则不带单位时不写单位"
    assert codec.format_sn(7, FULL, "51.5", "条") == "光-SN0007-051.5M条#尾", "米数补 0、分隔符、单位都写入"
    assert codec.format_sn(7, FULL, "123456", "条") == "光-SN0007-123456M条#尾", "米数超过补 0 位数不截断"
    assert codec.format_sn(7, UNIT_ONLY, "100", "PCS") == "U007PCS"
    assert codec.encode_seq(31, B32) == "000Z"
    assert codec.encode_seq(32, B32) == "0010"


def test_default_32_charset_drops_confusable_letters():
    assert DEFAULT_CHARSETS[32] == "0123456789ABCDEFGHJKMNPQRSTVWXYZ"
    assert not set("ILOU") & set(DEFAULT_CHARSETS[32])


def test_rejects_overflow_instead_of_widening():
    with pytest.raises(codec.SeqOverflow):
        codec.format_sn(100_000, NOLEN)
    assert codec.max_seq(NOLEN) == 99_999
    assert codec.max_seq(B32) == 32**4 - 1


def test_parse_recovers_serial():
    assert codec.parse("HX000001100A", HX) == 1
    assert codec.parse("HX00004210.5A", HX) == 42
    assert codec.parse("HX0000420A", HX) == 42
    assert codec.parse("HX000042000A", HX) == 42, "默认米数 000 原样写入"
    assert codec.parse("AVM00042X", NOLEN) == 42
    assert codec.parse("AVM0042X", NOLEN) is None, "wrong length"
    assert codec.parse("BVM00042X", NOLEN) is None, "wrong prefix"
    assert codec.parse("AVM00042100X", NOLEN) is None, "不带米数、单位的规则流水后不能有别的"
    assert codec.parse("光-SN0007-051.5M条#尾", FULL) == 7
    assert codec.parse("光-SN0007051.5M条#尾", FULL) is None, "米数前分隔符不符"
    assert codec.parse("U007PCS", UNIT_ONLY) == 7
    for s in (0, 1, 31, 32, 1023, 1_048_575):
        assert codec.parse(codec.format_sn(s, B32), B32) == s


def test_chinese_length_is_detected():
    for t in ("五十米", "51米", "５１", "51，5"):
        assert codec.has_chinese(t), t
    for t in ("51.5", "000", "L-2", "", None):
        assert not codec.has_chinese(t), t


@pytest.mark.parametrize(
    ("raw", "text"),
    [
        ("0100", "0100"),
        ("000", "000"),
        ("001", "001"),
        ("10.50", "10.50"),
        (" 100 ", "100"),
        ("L-2", "L-2"),
        ("", None),
        ("  ", None),
        (None, None),
    ],
)
def test_length_is_plain_string(raw, text):
    assert codec.clean_length(raw) == text


def test_validation():
    assert codec.validate(Spec("A", "", 8, 4, "01234567")) is ErrorCode.RULE_BASE_INVALID
    assert codec.validate(Spec("A", "", 10, 13, "0123456789")) is ErrorCode.RULE_SEQ_LEN_INVALID
    assert codec.validate(Spec("A", "", 10, 4, "012345678")) is ErrorCode.RULE_CHARSET_LENGTH
    assert codec.validate(Spec("A", "", 10, 4, "0123456788")) is ErrorCode.RULE_CHARSET_DUPLICATE
    assert codec.validate(Spec("", "", 10, 4, "0123456789")) is ErrorCode.RULE_PREFIX_REQUIRED
    for prefix in ("A_B", "ab", "A B", "A.B", "A/B", "A#", "SN😀", "前缀", "ＳＮ"):
        assert codec.validate(Spec(prefix, "x-1.#", 10, 4, "0123456789")) is None, "前缀、后缀不限制字符"
    assert codec.validate(Spec("A", "", 10, 4, "0123456789", True, 33)) is ErrorCode.RULE_LENGTH_PAD_INVALID
    assert codec.validate(Spec("A", "", 10, 4, "0123456789", True, 0, "-" * 17)) is ErrorCode.RULE_AFFIX_INVALID
    assert codec.validate(Spec("A", "", 10, 4, "01234 6789")) is ErrorCode.RULE_CHARSET_INVALID
    assert codec.validate(Spec("A" * 40, "B" * 10, 10, 4, "0123456789", True, 12)) is ErrorCode.RULE_SN_TOO_LONG
    assert codec.validate(Spec("A" * 40, "B" * 10, 10, 4, "0123456789", False, 12)) is None, "不带米数时补 0 不占长"
    assert codec.validate(HX) is None
    assert codec.validate(FULL) is None


def test_custom_rule_defaults_four_digits_without_length():
    from app.schemas.rule import CustomRuleIn
    from app.services import rules

    c = CustomRuleIn(prefix="6090001").model_dump()
    s = rules._custom_spec(c)
    assert (s.seq_len, s.with_length, s.with_unit) == (4, False, False)
    assert codec.format_sn(1, s, "100", "M") == "60900010001"


def test_rule_format_text_names_length_as_meters():
    from app.services import rules

    assert rules.format_text(FULL) == "光-SN + 4 位 10 进制流水 + - + 米数（补足 5 位） + M + 单位 + #尾"
    assert rules.format_text(HX) == "HX + 6 位 10 进制流水 + 米数 + A"
    assert "米标" not in rules.format_text(FULL)


def test_spec_without_length_drops_padding_and_separators():
    from app.services import rules

    s = rules.normalize(rules.SpecIn(" P ", "", 10, 4, None, False, 5, "-", "M", True))
    assert (s.prefix, s.length_pad, s.sep_before, s.sep_after, s.with_unit) == ("P", 0, "", "", True)


# ---------------------------------------------------------------------- 金蝶解析（附录 D.1）


def test_query_rows_trim_spaces_and_read_header_fields():
    row = [
        "VT1",
        "C",
        "0001 ",
        " ",
        "PO1 ",
        "110",
        [{"Key": 2052, "Value": "越南"}],
        " ",
        "9148363",
        "缆",
        "0100",
        10.0,
        2.0,
        "2026-10-01T08:00:00",
        " 条 ",
    ]
    r = k3._row_of_query(row)
    assert (r.bill_no, r.customer_code, r.pi, r.po, r.org_code, r.org_name, r.seq) == (
        "VT1",
        "0001",
        "",
        "PO1",
        "110",
        "越南",
        2,
    )
    assert r.attach_no == "" and r.length_raw == "0100"
    assert r.qty == Decimal("10.0")
    assert r.modify_time == datetime(2026, 10, 1, 8, 0, 0)
    assert r.unit == "条", "产品单位去首尾空格"


def test_query_field_keys_follow_interface_doc():
    """前 12 列与《生产订单查询接口》的 FieldKeys 一致，顺序不能变。"""
    assert ",".join(k3.DETAIL_KEYS[:12]) == (
        "FBillNo,FDocumentStatus,F_ora_kh.FNumber,F_HX_PI,F_HX_PO,FPrdOrgId.FNumber,FPrdOrgId.FName,"
        "FTERM,FMaterialId.FNumber,FMaterialId.FName,Flength,FQty"
    )
    assert k3.DETAIL_KEYS[12:] == ("FTreeEntity_FSeq", "FModifyDate", "FProductUnit")
    assert k3.ORDER_BY == "FBillNo ASC, FTreeEntity_FEntryId ASC"


def test_view_reads_document_status_and_lines():
    doc = {
        "BillNo": "VT2026090001",
        "DocumentStatus": "D",
        "ModifyDate": "2026-09-21T16:18:41.827",
        "F_HX_PI": "000001",
        "F_HX_PO": " ",
        "F_ora_Base": None,
        "PrdOrgId": {"Number": "110", "Name": [{"Key": 2052, "Value": "光纤连接（越南）有限公司"}]},
        "TreeEntity": [
            {
                "Seq": 1,
                "MaterialId": {"Number": "9148363", "Name": "x"},
                "Qty": 10.0,
                "Flength": " ",
                "FTERM": "0002",
                "FProductUnit": "PCS ",
            }
        ],
    }
    rows = k3.rows_from_view(doc)
    assert rows[0].doc_status == "D"
    assert rows[0].customer_code == "", "基础资料为 null 视为空"
    assert rows[0].po == "" and rows[0].length_raw == ""
    assert rows[0].attach_no == "0002"
    assert rows[0].unit == "PCS"
    assert rows[0].modify_time == datetime(2026, 9, 21, 16, 18, 41)


def test_locale_name_picks_simplified_chinese():
    assert k3.locale_name([{"Key": 1033, "Value": "Lock"}, {"Key": 2052, "Value": "门锁"}]) == "门锁"


def test_error_is_extracted():
    err = [[{"Result": {"ResponseStatus": {"IsSuccess": False, "Errors": [{"Message": "字段不存在"}]}}}]]
    assert k3.error_of(err) == "字段不存在"
    assert k3.error_of([["a"]]) is None


# ---------------------------------------------------------------------- MES 报文


def test_gzip_body_is_byte_identical_across_calls():
    raw = mes_client.canonical({"a": 1, "b": "中"})
    assert mes_client.gzip_bytes(raw) == mes_client.gzip_bytes(raw)
    assert gzip.decompress(mes_client.gzip_bytes(raw)) == raw


def test_mes_request_gunzip_is_capped():
    """MES 接口的 gzip 请求体：正常报文照常解压；解压后超过上限、不是合法 gzip、被截断的都返回 400，不解进内存。"""
    from app.api.routes import mes
    from app.core.exceptions import MesError

    body = json.dumps({"pi": "PI-1", "billNo": "MO-1"}).encode("utf-8")
    assert mes._gunzip(gzip.compress(body)) == body
    assert len(mes._gunzip(gzip.compress(b" " * mes.MAX_INFLATED))) == mes.MAX_INFLATED, "恰好到上限仍可解"
    bomb = gzip.compress(b"0" * (mes.MAX_INFLATED * 64))
    assert len(bomb) < 100 * 1024, "64 MB 的 0 压缩后不足 100 KB"
    for raw, text in ((bomb, "超过"), (b"not gzip", "gzip"), (gzip.compress(body)[:-6], "gzip")):
        with pytest.raises(MesError) as e:
            mes._gunzip(raw)
        assert (e.value.status_code, e.value.code) == (400, "INVALID_REQUEST") and text in e.value.message


def test_hmac_signature_string():
    import hashlib
    import hmac

    raw = b'{"x":1}'
    sig = mes_client.sign("k", "2026-10-19 10:37:00.123", raw)
    want = hmac.new(
        b"k", f"2026-10-19 10:37:00.123\n{hashlib.sha256(raw).hexdigest()}".encode(), hashlib.sha256
    ).hexdigest()
    assert sig == want


@pytest.mark.parametrize(
    ("auth", "want"),
    [("TOKEN", {"authorization"}), ("HMAC", {"x-sn-timestamp", "x-sn-signature"}), ("NONE", set())],
)
def test_auth_headers_by_type(auth, want):
    """无认证不带任何认证头；需要凭证的只有 TOKEN、HMAC。"""
    t = mes_client.Target("https://mes", auth, "k" if auth != "NONE" else "", 10)
    h = {k.lower() for k in mes_client._headers(t, 1, b"{}")}
    assert h & {"authorization", "x-sn-timestamp", "x-sn-signature"} == want
    assert mes_client.needs_credential(auth) is (auth != "NONE")


def test_ping_batch_has_no_items():
    body = json.loads(gzip.decompress(mes_client.ping_body("110", "PING-1")))
    assert set(body) == {
        "schemaVersion",
        "batchNo",
        "batchType",
        "orgCode",
        "toOrgCode",
        "count",
        "createdAt",
        "itemsDigest",
        "items",
    }
    assert body["batchType"] == "PING" and body["items"] == [] and body["toOrgCode"] is None and body["count"] == 0
    assert body["itemsDigest"].startswith("sha256:") and len(body["itemsDigest"]) == 71


@pytest.mark.parametrize("code,kind", [("OK", mes_client.RECEIVED), ("RECEIVED", mes_client.FAIL)])
def test_ping_only_accepts_received(code, kind):
    """连接测试与正式推送一样应答 OK 即收到；其他 code（如 RECEIVED）不算通过。"""
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from threading import Thread

    class H(BaseHTTPRequestHandler):
        def do_POST(self):
            self.rfile.read(int(self.headers["Content-Length"]))
            raw = json.dumps({"batchNo": "PING-1", "code": code, "message": ""}).encode()
            self.send_response(200)
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)

        def log_message(self, *a):
            pass

    srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
    Thread(target=srv.serve_forever, daemon=True).start()
    try:
        t = mes_client.Target(f"http://127.0.0.1:{srv.server_port}/", mes_client.NONE, "", 5)
        out = mes_client.send(t, "PING-1", mes_client.ping_body("110", "PING-1"), 1, ping=True)
    finally:
        srv.shutdown()
    assert out.kind == kind
    if kind == mes_client.FAIL:
        assert out.code == "RESPONSE_INVALID"


# ---------------------------------------------------------------------- 多语言契约（对应 I18nContractTest）

PH = re.compile(r"\{(\w+)}")


@pytest.mark.parametrize("lang", ["zh-CN", "en", "vi"])
def test_every_error_code_translated(lang):
    path = ROOT / "frontend" / "src" / "locales" / f"errors.{lang}.json"
    if not path.is_file():
        pytest.skip("frontend sources not present")
    msgs = json.loads(path.read_text(encoding="utf-8"))
    for c in ErrorCode:
        assert c.name in msgs, f"{lang} missing {c.name}"
        assert set(PH.findall(c.zh)) == set(PH.findall(msgs[c.name])), f"{lang} placeholders of {c.name}"
    assert len(msgs) == len(ErrorCode), f"{lang} has stale codes"


# ---------------------------------------------------------------------- 建表脚本


def test_migration_flyway_checksum():
    """测试阶段只保留一个 V1__init.sql；校验和按 Flyway 算法固定，改了建表脚本须同步更新这里（并重建测试库）。"""
    ours = {m.script: m for m in migrations()}
    assert list(ours) == ["V1__init.sql"]
    assert ours["V1__init.sql"].checksum() == -874897514
    statements = ours["V1__init.sql"].statements()
    assert len(statements) == 34
    text = "\n".join(statements)
    # 脚本独立完整：33 张表各一条 CREATE TABLE 加 sn_sync_state 初始行，不用 ALTER 等语句事后补字段、补索引
    kinds = [" ".join(s.split()[:2]).upper() for s in statements]
    assert kinds.count("CREATE TABLE") == 33 and kinds.count("INSERT INTO") == 1, kinds
    for table in (
        "sn_line_count",
        "sn_line_length",
        "sn_line_unit",
        "sn_delivery_batch",
        "sn_push_task",
        "sn_transfer_item",
    ):
        assert f"CREATE TABLE {table}" in text, table
    assert "ix_rule_factory (factory_code)" in text, "规则所属组织（V31 权限调整）"
    for gone in ("sn_print", "sn_pi_import", "print_no", "sn_allocation"):
        assert gone not in text, f"{gone} 已随 PRD V30 移除（不再打印、不再按 PI 导入）"
    # 开发阶段清库重建，没有旧数据：不留页面领取、旧协议逐枚回执、早期转组织单的字段
    for gone in ("sn_acquire_batch", "acquire_batch", "ACQUIRED", "SUCCESS_CONFLICT", "conflict_qty", "existing_rev"):
        assert gone not in text, gone
    assert "manual_fields" not in text


# ---------------------------------------------------------------------- 定时任务


def test_spring_cron_next_fire():
    c = Cron("0 10 2 * * *")
    assert c.next_after(datetime(2026, 9, 28, 1, 0, 0)) == datetime(2026, 9, 28, 2, 10, 0)
    assert c.next_after(datetime(2026, 9, 28, 2, 10, 0)) == datetime(2026, 9, 29, 2, 10, 0)
    assert Cron("0 0 3 * * SUN").next_after(datetime(2026, 9, 28)) == datetime(2026, 10, 4, 3, 0, 0)
    assert Cron("*/15 * * * * *").next_after(datetime(2026, 1, 1, 0, 0, 1)) == datetime(2026, 1, 1, 0, 0, 15)


# ---------------------------------------------------------------------- 导出防公式注入


def test_xlsx_writes_formula_like_text_as_plain_string():
    import io

    import openpyxl

    rows = [["=1+1", '=HYPERLINK("http://x","y")', "@x", "-AB0001", 7]]
    raw = files.to_bytes(files.xlsx_bytes("t", ["=h"], rows))
    ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    assert [c.value for c in cells] == ["=h", "=1+1", '=HYPERLINK("http://x","y")', "@x", "-AB0001", 7], (
        "原样，不加前缀"
    )
    assert all(c.data_type == "s" for c in cells[:-1])
    assert cells[-1].data_type == "n"


# ---------------------------------------------------------------------- 请求号


@pytest.mark.parametrize("raw", ["abc-123", "0f1e2d3c4b5a69788796a5b4c3d2e1f0", "a.b:c_d", "x" * 64])
def test_request_id_kept_when_valid(raw):
    assert request_id_of(raw) == raw


@pytest.mark.parametrize("raw", [None, "", "x" * 65, "a b", "id\r\nX: y", "中文"])
def test_request_id_regenerated_when_invalid(raw):
    rid = request_id_of(raw)
    assert rid != raw and len(rid) == 36


# ---------------------------------------------------------------------- 令牌签名密钥


@pytest.mark.parametrize(
    "secret,problem",
    [
        ("", True),
        ("change-me-sn-center-secret", True),
        ("change-me-to-a-long-random-string", True),
        ("short-secret", True),
        ("x" * 32, False),
    ],
)
def test_secret_problem(monkeypatch, secret, problem):
    from app.core.config import settings

    monkeypatch.setenv("SN_SECRET", secret)
    assert (settings.secret_problem() is not None) is problem


def test_production_refuses_weak_secret_but_local_only_warns(monkeypatch):
    from app.core.config import settings
    from app.core.security import _sign
    from app.main import check_secret

    monkeypatch.setenv("SN_SECRET", "")
    monkeypatch.setenv("ENVIRONMENT", "local")
    check_secret()
    assert settings.secret_key, "blank secret never signs with an empty key"
    assert _sign("x") == _sign("x") != ""
    monkeypatch.setenv("ENVIRONMENT", "production")
    with pytest.raises(RuntimeError):
        check_secret()
    monkeypatch.setenv("SN_SECRET", "change-me-to-a-long-random-string")
    with pytest.raises(RuntimeError):
        check_secret()
    monkeypatch.setenv("SN_SECRET", "0123456789abcdef0123456789abcdef")
    check_secret()


# ---------------------------------------------------------------------- 数据库连接配置

DB_KEYS = ("DB_URL", "DB_HOST", "DB_PORT", "DB_NAME", "DB_USER", "DB_PASSWORD")


def _db_target(monkeypatch, **dotenv: str):
    from app.core.config import DbTarget, Settings

    for k in DB_KEYS:
        monkeypatch.delenv(k, raising=False)
    s = Settings()
    s._dotenv = dict(dotenv)
    t = s.db_target()
    assert isinstance(t, DbTarget)
    return t


def test_db_target_reads_separate_keys(monkeypatch):
    t = _db_target(
        monkeypatch, DB_HOST="db.example", DB_PORT="3307", DB_NAME="other_db", DB_USER="u2", DB_PASSWORD=" p#w=d "
    )
    assert (t.host, t.port, t.database, t.user, t.password) == ("db.example", 3307, "other_db", "u2", " p#w=d ")


def test_db_target_defaults(monkeypatch):
    t = _db_target(monkeypatch)
    assert (t.host, t.port, t.database, t.user, t.password) == ("127.0.0.1", 3306, "sn_dev", "appuser", "")


def test_db_target_legacy_url_still_works_and_single_keys_win(monkeypatch):
    url = "jdbc:mysql://10.0.0.5:3310/legacy?useUnicode=true&characterEncoding=utf8"
    t = _db_target(monkeypatch, DB_URL=url, DB_USER="appuser", DB_PASSWORD="x")
    assert (t.host, t.port, t.database) == ("10.0.0.5", 3310, "legacy")
    t = _db_target(monkeypatch, DB_URL=url, DB_NAME="sn_test", DB_HOST="10.0.0.6")
    assert (t.host, t.port, t.database) == ("10.0.0.6", 3310, "sn_test")
    t = _db_target(monkeypatch, DB_URL="mysql://me:p%40ss@h1/d1")
    assert (t.host, t.port, t.database, t.user, t.password) == ("h1", 3306, "d1", "me", "p@ss")


def test_db_target_env_overrides_dotenv(monkeypatch):
    from app.core.config import Settings

    for k in DB_KEYS:
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("DB_NAME", "from_env")
    s = Settings()
    s._dotenv = {"DB_NAME": "from_file"}
    assert s.db_target().database == "from_env"


# ---------------------------------------------------------------------- 失败痕迹的动作归类


@pytest.mark.parametrize(
    ("method", "path", "action"),
    [
        ("POST", "/api/transfers/org/direct", "TRANSFER_DIRECT"),
        ("POST", "/api/transfers/bill", "TRANSFER_BILL"),
        ("POST", "/api/allocation/allocate", "SN_ALLOCATE"),
        ("POST", "/api/generate/run", "SN_GENERATE"),
        ("POST", "/api/generate/batches/G1/cancel", "SN_GEN_CANCEL"),
        ("POST", "/api/generate/batches/G1/revoke", "SN_REVOKE"),
        ("PUT", "/api/rules/pi-binding", "RULE_PI_BIND"),
        ("DELETE", "/api/rules/pi-binding", "RULE_PI_UNBIND"),
        ("POST", "/api/rules/line-select", "RULE_LINE_SELECT"),
        ("DELETE", "/api/rules/3", "RULE_DELETE"),
        ("PUT", "/api/rules/3", "RULE_UPDATE"),
        ("POST", "/api/users", "USER_CREATE"),
        ("PUT", "/api/users/5", "USER_UPDATE"),
        ("POST", "/api/users/5/disable", "USER_DISABLE"),
        ("POST", "/api/users/5/enable", "USER_ENABLE"),
        ("POST", "/api/users/5/reset-password", "USER_RESET_PWD"),
        ("POST", "/api/users/5/reset-secret", "USER_RESET_SECRET"),
        ("PUT", "/api/factories/110/mes-config", "FACTORY_MES_CONFIG"),
        ("POST", "/api/factories/110/mes-toggle", "FACTORY_MES_TOGGLE"),
        ("POST", "/api/delivery/batches/PS-1/void", "SN_BATCH_VOID"),
        ("POST", "/api/delivery/batches/PS-1/confirm", "SN_BATCH_MANUAL"),
        ("POST", "/api/delivery/resume", "SN_PUSH_RESUME"),
        ("POST", "/api/orders/attach", "ORDER_ATTACH_FILL"),
        ("PUT", "/api/orders/length", "ORDER_LENGTH_FILL"),
        ("PUT", "/api/orders/unit", "ORDER_UNIT_FILL"),
        ("POST", "/api/generate/revoke-config", "LINE_CONFIG_REVOKE"),
        ("POST", "/api/generate/revoke", "SN_REVOKE"),
        ("POST", "/api/orders/alerts/3/handle", "ORDER_ALERT_HANDLE"),
        ("POST", "/api/mes/v1/batches/next", "SN_PULL"),
        ("POST", "/api/mes/v1/receipts", "SN_RECEIPT"),
        ("PUT", "/api/auth/lang", "OTHER"),
    ],
)
def test_failed_write_action_follows_the_path(method, path, action):
    from app.services.audit import action_for

    assert action_for(method, path) == action
