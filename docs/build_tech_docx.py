"""把 docs/系统技术方案.md 转成 docs/系统技术方案.docx。

1. 用 pandoc 转 docx（中文字体模板），md 中引用的图片按页宽缩放、居中，图题居中；
2. 生成带超链接和页码的目录（Word 目录域，打开后右键「更新域」可刷新），目录后分页；
3. 表格加边框、表头底色，页脚加页码；
4. 装有 LibreOffice 时先转 PDF 算出各标题所在页，写进目录；没有则目录只有标题、没有页码。

依赖：pip install pypandoc_binary python-docx pillow pymupdf；可选 LibreOffice（soffice）
用法（仓库根目录）：python docs/build_tech_docx.py
"""
import copy
import os
import re
import shutil
import subprocess
import tempfile

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

SRC = 'docs/系统技术方案.md'
OUT = 'docs/系统技术方案.docx'
TMP = tempfile.mkdtemp()
MAXW, MAXH = 16.0, 21.0  # 图片最大宽 / 高（cm）


def prepare_markdown():
    s = open(SRC, encoding='utf-8').read()
    s = re.sub(r'## 目录\n.*?(?=\n## 第一章)', '', s, flags=re.S)  # 手写目录由 Word 目录代替
    title = s.split('\n', 1)[0].lstrip('# ').strip()
    s = s.split('\n', 1)[1]

    def img(m):
        alt, path = m.group(1), 'docs/' + m.group(2)
        w, h = Image.open(path).size
        cw = MAXW
        if cw * h / w > MAXH:
            cw = MAXH * w / h
        return f'![]({path}){{width={cw:.1f}cm}}'
    s = re.sub(r'!\[([^\]]*)\]\((images/[^)]+)\)', img, s)
    s = re.sub(r'^\*\*(图 \d-\d[^*]*)\*\*$', r'::: {custom-style="Figure Caption"}\n\1\n:::', s, flags=re.M)
    md = f'{TMP}/doc.md'
    open(md, 'w', encoding='utf-8').write(f'---\ntitle: "{title}"\n---\n' + s)
    return md


def set_font(d, style, east, west, size=None, bold=None, color=None):
    st = d.styles[style]
    st.font.name = west
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(k), west)
    rf.set(qn('w:eastAsia'), east)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]
    if size:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if color is not None:
        st.font.color.rgb = RGBColor.from_string(color)


def reference_doc():
    ref = f'{TMP}/ref.docx'
    subprocess.run([pypandoc.get_pandoc_path(), '-o', ref, '--print-default-data-file', 'reference.docx'], check=True)
    d = Document(ref)
    for st in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        set_font(d, st, '宋体', 'Calibri', 10.5)
    set_font(d, 'Title', '黑体', 'Arial', 22, True, '1F3864')
    for i, sz in ((1, 16), (2, 14), (3, 12), (4, 11)):
        set_font(d, f'Heading {i}', '黑体', 'Arial', sz, True, '1F3864')
    set_font(d, 'Verbatim Char', '等线', 'Consolas', 9)
    d.save(ref)
    return ref


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def insert_ordered(parent, child, later):
    """按 schema 顺序插入：放在第一个属于 later 的子元素之前"""
    anchor = next((c for c in parent if c.tag in [qn(t) for t in later]), None)
    if anchor is None:
        parent.append(child)
    else:
        anchor.addprevious(child)


def text_len(text):
    return sum(2 if ord(ch) > 0x2e80 else 1 for ch in text)


def fit_columns(t, total=int(16.6 * 567)):
    """按各列内容长度分配列宽（twips），短列不被挤窄，长说明列拿到更多宽度"""
    n = len(t.columns)
    need = [6] * n
    for row in t.rows:
        for k, c in enumerate(row.cells[:n]):
            need[k] = max(need[k], min(text_len(c.text.strip()), 70))
    # 每列至少放得下最长的一个「词」（编码、表名不能被截断），剩余宽度按内容长度分
    floor = [0] * n
    for row in t.rows:
        for k, c in enumerate(row.cells[:n]):
            txt = c.text.strip()
            toks = [txt] if len(txt) <= 12 else re.split(r'[\s/，、；（）()]+', txt)
            for tok in toks:
                floor[k] = max(floor[k], min(len(tok) * 115 + 220, total // 3) if tok.isascii() else 0)
    weights = [max(w, 8) ** 0.8 for w in need]
    widths = [int(total * w / sum(weights)) for w in weights]
    short = [k for k in range(n) if widths[k] < floor[k]]
    if short:
        fixed = sum(floor[k] for k in short)
        others = [k for k in range(n) if k not in short]
        ow = sum(weights[k] for k in others) or 1
        for k in short:
            widths[k] = floor[k]
        for k in others:
            widths[k] = int((total - fixed) * weights[k] / ow)
    grid = t._tbl.tblGrid
    for k, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(widths[k]))
    for row in t.rows:
        for k, c in enumerate(row.cells[:n]):
            pr = c._tc.get_or_add_tcPr()
            for old in pr.findall(qn('w:tcW')):
                pr.remove(old)
            pr.insert(0, el('w:tcW', **{'w:w': widths[k], 'w:type': 'dxa'}))
    pr = t._tbl.tblPr
    for old in pr.findall(qn('w:tblLayout')):
        pr.remove(old)
    insert_ordered(pr, el('w:tblLayout', **{'w:type': 'fixed'}), ['w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription'])


def style_tables(d):
    for t in d.tables:
        pr = t._tbl.tblPr
        for tag in ('w:tblW', 'w:tblBorders'):
            for old in pr.findall(qn(tag)):
                pr.remove(old)
        borders = el('w:tblBorders')
        for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(el(f'w:{e}', **{'w:val': 'single', 'w:sz': 4, 'w:space': 0, 'w:color': '999999'}))
        later = ['w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription']
        insert_ordered(pr, el('w:tblW', **{'w:w': int(16.6 * 567), 'w:type': 'dxa'}), later + ['w:jc', 'w:tblCellSpacing', 'w:tblInd'])
        insert_ordered(pr, borders, later)
        fit_columns(t)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(el('w:tblHeader'))  # 跨页时重复表头
        for c in t.rows[0].cells:
            insert_ordered(c._tc.get_or_add_tcPr(), el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': 'D9E2F3'}),
                           ['w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark'])
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(9)


def headings(d):
    """[(级别, 文本, 书签名, 段落)]，只取标题 1、2"""
    out = []
    for p in d.paragraphs:
        name = p.style.name
        if name in ('Heading 1', 'Heading 2'):
            bm = p._p.find(qn('w:bookmarkStart'))
            if bm is None:
                bm_name = f'_toc{len(out)}'
                p._p.insert(1, el('w:bookmarkStart', **{'w:id': 9000 + len(out), 'w:name': bm_name}))
                p._p.append(el('w:bookmarkEnd', **{'w:id': 9000 + len(out)}))
            else:
                bm_name = bm.get(qn('w:name'))
            out.append((int(name[-1]), p.text.strip(), bm_name, p))
    return out


def run(text=None, **rpr):
    r = el('w:r')
    if rpr:
        pr = el('w:rPr')
        if rpr.get('bold'):
            pr.append(el('w:b'))
        if rpr.get('noproof'):
            pr.append(el('w:noProof'))
        r.append(pr)
    if text is not None:
        t = el('w:t')
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        t.text = text
        r.append(t)
    return r


def fld(kind, **attrs):
    r = el('w:r')
    r.append(el('w:fldChar', **{'w:fldCharType': kind, **attrs}))
    return r


def instr(text):
    r = el('w:r')
    t = el('w:instrText')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    r.append(t)
    return r


def toc_paragraphs(hs, pages, text_width_twips):
    """目录域：第一段带域开始，最后一段带域结束；每条是指向标题书签的超链接 + 点线 + 页码"""
    paras = []
    for k, (lvl, text, bm, _) in enumerate(hs):
        p = el('w:p')
        ppr = el('w:pPr')
        ppr.append(el('w:pStyle', **{'w:val': f'TOC{lvl}'}))
        tabs = el('w:tabs')
        tabs.append(el('w:tab', **{'w:val': 'right', 'w:leader': 'dot', 'w:pos': text_width_twips}))
        ppr.append(tabs)
        ppr.append(el('w:spacing', **{'w:before': 120 if lvl == 1 else 0, 'w:after': 40}))
        ppr.append(el('w:ind', **{'w:left': 0 if lvl == 1 else 420}))
        p.append(ppr)
        if k == 0:
            p.append(fld('begin', **{'w:dirty': 'true'}))
            p.append(instr(' TOC \\o "1-2" \\h \\z \\u '))
            p.append(fld('separate'))
        h = el('w:hyperlink', **{'w:anchor': bm, 'w:history': 1})
        h.append(run(text, bold=(lvl == 1), noproof=True))
        h.append(run(None))
        h[-1].append(el('w:tab'))
        h.append(fld('begin'))
        h.append(instr(f' PAGEREF {bm} \\h '))
        h.append(fld('separate'))
        h.append(run(str(pages.get(bm, '')), noproof=True))
        h.append(fld('end'))
        p.append(h)
        if k == len(hs) - 1:
            p.append(fld('end'))
        paras.append(p)
    return paras


def page_footer(d):
    ft = d.sections[0].footer
    p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for x in (run('第 '), fld('begin'), instr(' PAGE '), fld('separate'), run('1'), fld('end'), run(' 页')):
        p._p.append(x)


def build(pages):
    ref = reference_doc()
    pypandoc.convert_file(prepare_markdown(), 'docx', outputfile=OUT,
                          extra_args=[f'--reference-doc={ref}', '--shift-heading-level-by=-1', '--resource-path=.'])
    d = Document(OUT)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.2)
    sec.gutter = Cm(0)
    style_tables(d)
    for p in d.paragraphs:
        if p.style.name == 'Figure Caption':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
        if any(x.tag == qn('w:drawing') for r in p.runs for x in r._r):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
    hs = headings(d)
    # 目录插在标题之后、正文之前，目录后分页
    title = next(p for p in d.paragraphs if p.style.name == 'Title')
    head = copy.deepcopy(title._p)
    for r in head.findall(qn('w:r')):
        head.remove(r)
    head.find(qn('w:pPr')).find(qn('w:pStyle')).set(qn('w:val'), 'TOCHeading')
    head.append(run('目　录', bold=True))
    width = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU → twips
    anchor = title._p
    for x in [head] + toc_paragraphs(hs, pages, width):
        anchor.addnext(x)
        anchor = x
    brk = el('w:p')
    brk.append(run(None))
    brk[-1].append(el('w:br', **{'w:type': 'page'}))
    anchor.addnext(brk)
    # 目录样式：TOC 1 / TOC 2 可能不在模板里，缺了就补
    for i in (1, 2):
        sid = f'TOC{i}'
        if not any(x.style_id == sid for x in d.styles):
            st = d.styles.add_style(f'toc {i}', 1)
            st.element.set(qn('w:styleId'), sid)
            st.base_style = d.styles['Normal']
    set_font(d, 'TOC Heading', '黑体', 'Arial', 16, True, '1F3864')
    page_footer(d)
    d.save(OUT)
    return [(bm, text) for _, text, bm, _ in hs]


def pdf_pages(hs):
    """用 LibreOffice 转 PDF，按标题文本找页码"""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return {}
    import pymupdf
    src = f'{TMP}/doc.docx'
    shutil.copy(OUT, src)
    subprocess.run([soffice, '--headless', '--norestore', f'-env:UserInstallation=file://{TMP}/lo',
                    '--convert-to', 'pdf', '--outdir', TMP, src], check=True, capture_output=True, timeout=300)
    doc = pymupdf.open(f'{TMP}/doc.pdf')
    texts = [pg.get_text().replace(' ', '') for pg in doc]
    pages, start = {}, 0
    toc_end = next((i for i, t in enumerate(texts) if '目录' in t.replace('　', '')), 0) + 1
    for bm, text in hs:
        key = text.replace(' ', '')
        for i in range(max(start, toc_end), len(texts)):
            if key in texts[i]:
                pages[bm] = i + 1
                start = i
                break
    return pages


if __name__ == '__main__':
    hs = build({})
    pages = pdf_pages(hs)
    if pages:
        build(pages)  # 目录条数不变，页码写进去后版面不变
    print('ok', f'{len(pages)}/{len(hs)} 个标题有页码')
