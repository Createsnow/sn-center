"""Read an actual browser download without changing its bytes or column order."""

import csv
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

path = Path(sys.argv[1]).resolve(strict=True)
if path.suffix == ".csv":
    with path.open(encoding="utf-8-sig", newline="") as stream:
        raw = list(csv.reader(stream))
    types = [["text"] * len(row) for row in raw]
    sheet = None
elif path.suffix == ".xlsx":
    book = load_workbook(path, read_only=True)
    sheet = book.active.title
    cells = list(book.active.iter_rows())
    raw = [[cell.value for cell in row] for row in cells]
    types = [[cell.data_type for cell in row] for row in cells]
    book.close()
else:
    raise ValueError("Expected browser CSV/XLSX download")
print(json.dumps({
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "sheet": sheet,
    "raw": raw,
    "types": types,
    "rows": [["" if value is None else str(value) for value in row] for row in raw],
}, ensure_ascii=False))
